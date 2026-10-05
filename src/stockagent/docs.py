"""文档与手动数据：文件入库与文本抽取、收件箱导入、手动数据 CSV、定期报告自动下载。"""

from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
from datetime import timedelta
from pathlib import Path

from .config import Company, Watchlist
from .htmltext import html_text  # noqa: F401  （docs.html_text 保持可用）
from .sources import cninfo, hkex, ids, sec
from .utils import iso, log, sha256_file, to_float, today_local, utcnow

DOC_TYPES = {
    "annual_report": "年报",
    "interim_report": "中期报告",
    "quarterly_report": "季报",
    "prospectus": "招股书",
    "research": "研报",
    "transcript": "纪要/电话会",
    "note": "笔记",
    "other": "其他",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".html", ".htm", *IMAGE_SUFFIXES}
OCR_LANGS = "chi_sim+eng"

# 顺序有意义："2026年半年度报告"同时含"半年"和"年度报告"，应判为中期报告
_TYPE_RULES = [
    (re.compile(r"招股|prospectus", re.I), "prospectus"),
    (re.compile(r"中期|半年|interim|half[-_ ]?year|\bH1\b", re.I), "interim_report"),
    (re.compile(r"季报|季報|季度|quarter|\bQ[1-4]\b", re.I), "quarterly_report"),
    (re.compile(r"年报|年報|年度报告|年度報告|annual", re.I), "annual_report"),
    (re.compile(r"纪要|紀要|调研|調研|电话会|電話會|transcript", re.I), "transcript"),
    (re.compile(r"研报|研報|深度|点评|點評|证券|證券|research", re.I), "research"),
    (re.compile(r"笔记|筆記|note", re.I), "note"),
]

_CN_DIGITS = str.maketrans("〇零一二三四五六七八九", "00123456789")


def infer_doc_type(name: str) -> str:
    for pattern, doc_type in _TYPE_RULES:
        if pattern.search(name):
            return doc_type
    return "other"


def infer_period(name: str, doc_type: str) -> str | None:
    # "二零二六年中期報告" → "2026年中期報告"
    name = re.sub(r"[〇零一二三四五六七八九]{4}", lambda m: m.group(0).translate(_CN_DIGITS), name)
    m = re.search(r"(20\d{2})", name)
    if not m:
        return None
    year = m.group(1)
    if doc_type == "interim_report":
        return f"{year}H1"
    if doc_type == "quarterly_report":
        q = re.search(r"Q([1-4])|第([一二三四1-4])季|([一三])季", name, re.I)
        if q:
            n = next(g for g in q.groups() if g)
            return f"{year}Q{'一二三四'.index(n) + 1 if n in '一二三四' else n}"
    return year


def _read_text(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def ocr_image(path: Path) -> str:
    """本地 tesseract 识别图片文字（固定参数调用，不经 shell）。"""
    exe = shutil.which("tesseract")
    if not exe:
        raise RuntimeError("未安装 tesseract，无法识别图片文字（brew install tesseract tesseract-lang）")
    p = subprocess.run([exe, str(path), "-", "-l", OCR_LANGS], capture_output=True, timeout=180)
    if p.returncode != 0:
        raise RuntimeError(f"tesseract 识别失败：{p.stderr.decode('utf-8', 'replace')[:200]}")
    return p.stdout.decode("utf-8", "replace")


def extract_text(path: Path) -> tuple[str, int | None]:
    """返回 (文本, 页数)。PDF 每页前插入 [第N页] 标记，方便模型引用出处；图片用 OCR。"""
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        return ocr_image(path), None
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(path) as doc:
            parts = []
            for i, page in enumerate(doc, 1):
                t = page.get_text("text").strip()
                if t:
                    parts.append(f"[第{i}页]\n{t}")
            return "\n\n".join(parts), doc.page_count
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines), None
    if suffix in (".html", ".htm"):
        return html_text(_read_text(path))
    return _read_text(path), None


def store_document(db, data_dir: Path, company_id: str, src: Path, *, doc_type: str, period: str | None,
                   title: str, source: str, source_url: str | None = None,
                   published_at: str | None = None) -> tuple[dict, bool]:
    """复制到 data/docs/<公司>/ 并抽取文本。按文件哈希去重，返回 (文档记录, 是否新增)。"""
    sha = sha256_file(src)
    existing = db.doc_by_sha(sha)
    if existing:
        return existing, False
    dest_dir = data_dir / "docs" / company_id
    dest_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w.\-]+", "_", src.stem)[:60] or "doc"
    dest = dest_dir / f"{sha[:8]}_{safe}{src.suffix.lower()}"
    if src.resolve() != dest.resolve():
        shutil.copy2(src, dest)
    text, pages = extract_text(dest)
    text_path = dest.with_name(dest.name + ".txt")
    text_path.write_text(text, encoding="utf-8")
    doc = {
        "company_id": company_id, "doc_type": doc_type, "period": period, "title": title,
        "source": source, "source_url": source_url, "file_path": str(dest), "text_path": str(text_path),
        "sha256": sha, "pages": pages, "chars": len(text), "published_at": published_at,
    }
    doc["id"] = db.add_document(doc)
    if len(text) < 1000 and dest.suffix.lower() == ".pdf":
        log.warning("%s 几乎没有可抽取的文字（可能是扫描件，需要 OCR 后再导入）", src.name)
    return doc, True


def read_document_text(doc: dict) -> str:
    return Path(doc["text_path"]).read_text(encoding="utf-8")


# ---------------------------------------------------------------- 手动数据

_HEADER_ALIASES = {
    "指标": "metric", "期间": "period", "日期": "period", "数值": "value", "值": "value",
    "单位": "unit", "备注": "note", "来源": "source", "公司": "company",
}


def import_manual_csv(db, path: Path, company_id: str | None, watchlist: Watchlist) -> str:
    """列：metric, period, value[, unit, note, source, company]（也接受中文列名）。
    company 为空表示行业/宏观指标（如 R32 制冷剂价格）。"""
    rows = list(csv.DictReader(io.StringIO(_read_text(path))))
    n = 0
    for raw in rows:
        row = {_HEADER_ALIASES.get((k or "").strip(), (k or "").strip().lower()): (v or "").strip() for k, v in raw.items()}
        if not row.get("metric") or not row.get("period"):
            continue
        cid = company_id or ""
        if not cid and row.get("company"):
            c = watchlist.get(row["company"])
            if c is None:
                raise ValueError(f"{path.name}: 找不到公司 {row['company']}")
            cid = c.id
        value = to_float((row.get("value") or "").replace(",", ""))
        db.upsert_manual(cid, row["metric"], row["period"], value, row.get("unit") or None,
                         row.get("note") or None, row.get("source") or f"manual:{path.name}")
        n += 1
    return f"导入 {n} 条手动数据 ← {path.name}"


# ---------------------------------------------------------------- 收件箱

def _archive(path: Path, dest_dir: Path) -> None:
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / path.name
    i = 1
    while dest.exists():
        dest = dest_dir / f"{path.stem}_{i}{path.suffix}"
        i += 1
    shutil.move(str(path), dest)


def import_inbox(db, data_dir: Path, watchlist: Watchlist) -> list[str]:
    """data/inbox/<公司id或名称或代码>/ 下的文件入库；inbox 根目录下的 CSV 视为手动数据。
    成功导入的文件移到 inbox/_imported/<日期>/，失败的留在原处。"""
    inbox = data_dir / "inbox"
    inbox.mkdir(parents=True, exist_ok=True)
    archive = inbox / "_imported" / today_local().isoformat()
    msgs: list[str] = []
    for entry in sorted(inbox.iterdir()):
        if entry.name.startswith((".", "_")):
            continue
        if entry.is_file():
            if entry.suffix.lower() == ".csv":
                try:
                    msgs.append(import_manual_csv(db, entry, None, watchlist))
                    _archive(entry, archive)
                except Exception as e:
                    msgs.append(f"✗ {entry.name}: {e}")
            else:
                msgs.append(f"跳过 {entry.name}：文档请放进公司子目录，例如 inbox/{watchlist.companies[0].id}/")
            continue
        company = watchlist.get(entry.name)
        if company is None:
            msgs.append(f"跳过目录 {entry.name}：不在自选股中（目录名可用公司 id、名称或代码）")
            continue
        for f in sorted(p for p in entry.rglob("*") if p.is_file() and not p.name.startswith(".")):
            try:
                if f.suffix.lower() == ".csv":
                    msgs.append(import_manual_csv(db, f, company.id, watchlist))
                elif f.suffix.lower() in SUPPORTED:
                    doc_type = infer_doc_type(f.name)
                    doc, new = store_document(
                        db, data_dir, company.id, f, doc_type=doc_type, period=infer_period(f.name, doc_type),
                        title=f.stem, source="manual",
                    )
                    msgs.append(
                        f"{'导入' if new else '已存在'}：{company.name} · {DOC_TYPES[doc['doc_type']]} "
                        f"{doc['period'] or ''} · {f.name}（{doc['chars']:,} 字）"
                    )
                else:
                    msgs.append(f"跳过 {f.name}：不支持的格式（支持 {', '.join(sorted(SUPPORTED))}, .csv）")
                    continue
                _archive(f, archive / entry.name)
            except Exception as e:
                msgs.append(f"✗ {f.name}: {e}")
    return msgs


# ---------------------------------------------------------------- 定期报告自动下载

def report_candidates(db, session, company: Company, days: int = 500, sec_user_agent: str | None = None) -> list[dict]:
    """最新年报/中期报告列表。A+H 公司优先用巨潮的 A 股简体全文；只在美股上市的公司取 SEC 的 10-K / 20-F / 10-Q。"""
    end, start = utcnow(), utcnow() - timedelta(days=days)
    a, h, u = company.listing("A"), company.listing("HK"), company.listing("US")
    if a:
        org = ids.cninfo_org_id(db, session, a)
        if org:
            found = cninfo.reports(session, a, org, start, end)
            if found:
                return [{**x, "source": "cninfo"} for x in found]
    if h:
        sid = ids.hkex_stock_id(db, session, h)
        if sid:
            return [{**x, "source": "hkex"} for x in hkex.reports(session, sid, start, end)]
    if u and not a and not h and sec_user_agent:
        cik = ids.sec_cik(db, session, u, sec_user_agent)
        if cik:
            return [{**x, "source": "sec"} for x in sec.reports(session, cik, sec_user_agent, start, end)]
    return []


def fetch_reports(db, session, data_dir: Path, company: Company, sec_user_agent: str | None = None) -> list[str]:
    msgs = []
    us_only = bool(company.listing("US")) and not company.listing("A") and not company.listing("HK")
    candidates = report_candidates(db, session, company, days=1200 if us_only else 500, sec_user_agent=sec_user_agent)
    if us_only and not sec_user_agent:
        return [f"{company.name}：跳过美股定期报告，config.yaml 未设置 sources.sec_user_agent"]
    # 每类取最新几份：美股公司年报取 3 份（看分部、客户集中度等的多年趋势），另取最新一份 10-Q 与代理声明（DEF 14A，记为「其他」）
    want = {"annual_report": 3 if us_only else 1, "interim_report": 1, "quarterly_report": 1, "other": 1}
    for doc_type, n in want.items():
        picks = [c for c in candidates if c["doc_type"] == doc_type][:n]
        if doc_type in ("quarterly_report", "other") and not picks:
            continue
        if doc_type == "interim_report" and us_only:
            continue
        if not picks:
            msgs.append(f"{company.name}：未找到{DOC_TYPES[doc_type]}")
            continue
        for latest in picks:
            period = latest.get("period") or infer_period(latest["title"], doc_type)
            if db.has_document(company.id, doc_type, period):
                continue
            tmp_dir = data_dir / "tmp"
            tmp_dir.mkdir(parents=True, exist_ok=True)
            suffix = Path(latest["url"].split("?")[0]).suffix.lower()
            suffix = suffix if suffix in (".htm", ".html") else ".pdf"
            tmp = tmp_dir / (re.sub(r"[^\w.\-]+", "_", f"{company.name}_{latest['title']}")[:80] + suffix)
            headers = {"User-Agent": sec_user_agent} if latest["source"] == "sec" else None  # SEC 要求带联系方式的 User-Agent
            try:
                with session.get(latest["url"], headers=headers, stream=True, timeout=180) as r:
                    r.raise_for_status()
                    with open(tmp, "wb") as f:
                        for chunk in r.iter_content(1 << 16):
                            f.write(chunk)
                doc, _ = store_document(
                    db, data_dir, company.id, tmp, doc_type=doc_type, period=period, title=latest["title"],
                    source=latest["source"], source_url=latest["url"], published_at=iso(latest["published_at"]),
                )
                msgs.append(f"下载：{company.name} · {latest['title']}（{doc['pages'] or '—'} 页，{doc['chars']:,} 字）")
            except Exception as e:
                msgs.append(f"✗ 下载 {company.name} {latest['title']} 失败：{e}")
            finally:
                tmp.unlink(missing_ok=True)
    return msgs
